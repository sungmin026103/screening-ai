"""SR Studio V36 — Supplementary Tables 생성(Excel · Word).

메타분석 결과 객체(auto_figures.analyze_outcome / meta_sections.result_from_*)에서
논문 Supplementary용 표 S1–S9를 만든다. 숫자는 figure와 같은 계산 경로(meta_sections)에서 나온다.

  S1  Effect sizes included in each meta-analysis (effect level)
  S2  Pooled estimates from three-level random-effects models
  S3  Subgroup analyses (Q_M omnibus test)
  S4  Leave-one-study-out sensitivity analysis
  S5  Influence and outlier diagnostics
  S6  Robustness of pooled estimates (sensitivity summary)
  S7  Small-study effects and publication bias (Egger, trim-and-fill)
  S8  Meta-regression, dose–response and advanced bias analyses
  S9  Data extraction quality control
"""
from __future__ import annotations

import io
from dataclasses import dataclass

import numpy as np
import pandas as pd

import meta_sections as M

ABBR = ("CI, confidence interval; CR2, bias-reduced cluster-robust variance estimator (Satterthwaite df); "
        "g, Hedges' g (standardized mean difference); k, number of effect sizes; PI, prediction interval; "
        "REML, restricted maximum likelihood; HKSJ, Hartung–Knapp–Sidik–Jonkman; "
        "τ²(L2), within-study variance; τ²(L3), between-study variance.")


@dataclass
class SuppTable:
    code: str
    title: str
    df: pd.DataFrame
    note: str = ""


def _p(p) -> str:
    try:
        p = float(p)
    except (TypeError, ValueError):
        return ""
    if not np.isfinite(p):
        return ""
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def _f(x, d=2) -> str:
    try:
        x = float(x)
    except (TypeError, ValueError):
        return "" if x is None else str(x)
    return "" if not np.isfinite(x) else f"{x:.{d}f}"


def _ci(lo, hi, d=2) -> str:
    a, b = _f(lo, d), _f(hi, d)
    return f"{a} to {b}" if a and b else ""


def _msd(m, s) -> str:
    def g3(v):
        try:
            v = float(v)
        except (TypeError, ValueError):
            return ""
        if not np.isfinite(v):
            return ""
        return f"{v:.3g}" if v != 0 and abs(v) < 10 else f"{v:.2f}"
    a, b = g3(m), g3(s)
    return f"{a} ± {b}" if a and b else ""


def _vardecomp(res: dict) -> tuple[float, float, float]:
    vi = np.asarray(res["vi"], float)
    k = len(vi)
    f = res["fit"]
    w = 1 / vi
    vt = (w.sum() * (k - 1)) / (w.sum() ** 2 - (w ** 2).sum()) if k > 1 else np.nan
    tot = vt + f.tau2_L2 + f.tau2_L3
    return 100 * vt / tot, 100 * f.tau2_L2 / tot, 100 * f.tau2_L3 / tot


def table_s1(results: list[dict], label_mode: str = "plain") -> SuppTable:
    rows = []
    for r in results:
        d, sub = r["data"], r["sub"]
        for i in range(len(d)):
            row = d.iloc[i]
            rows.append({
                "Outcome": r["outcome"],
                "Study": sub["study_plain"].iloc[i] if label_mode == "plain" else sub["study"].iloc[i],
                "Intervention": row.get("Intervention", ""),
                "Dose": row.get("dose", ""),
                "Species/model": row.get("Species", ""),
                "Duration (d)": _f(row.get("Intervention_day", np.nan), 0),
                "n (I/C)": (f"{_f(row.get('N_treat'), 0)}/{_f(row.get('N_control'), 0)}"
                            if "N_treat" in d.columns else ""),
                "Intervention mean ± SD": _msd(row.get("Mean_treat"), row.get("SD_treat")),
                "Control mean ± SD": _msd(row.get("Mean_control"), row.get("SD_control")),
                "Hedges' g": _f(sub["yi"].iloc[i]),
                "SE": _f(np.sqrt(sub["vi"].iloc[i]), 3),
                "95% CI": _ci(sub["ci_lo"].iloc[i], sub["ci_hi"].iloc[i]),
                "Weight (%)": _f(sub["weight_pct"].iloc[i], 1),
            })
    df = pd.DataFrame(rows)
    df = df.loc[:, [c for c in df.columns if df[c].astype(str).str.strip().ne("").any()]]
    note = ("Effect sizes are Hedges' g with the exact small-sample correction (metafor::escalc, measure = 'SMD'); "
            "weights are from the three-level model (Study/effect). Negative g indicates a lower value in the "
            "intervention group than in the control group.")
    return SuppTable("S1", "Effect sizes included in each meta-analysis", df, note)


def table_s2(results: list[dict]) -> SuppTable:
    rows = []
    for r in results:
        f = r["fit"]
        ps, p2, p3 = _vardecomp(r)
        rows.append({
            "Outcome": r["outcome"], "k": f.k, "Studies": f.n_studies, "Hedges' g": _f(f.mu),
            "95% CI (CR2)": _ci(f.cr2_ci_lb, f.cr2_ci_ub), "p (CR2)": _p(f.cr2_p), "CR2 df": _f(f.cr2_df, 1),
            "95% CI (model)": _ci(f.ci_lb, f.ci_ub), "p (model)": _p(f.pval),
            "95% PI": _ci(f.pi_lb, f.pi_ub), "I² (%)": _f(f.i2, 1),
            "τ²(L2)": _f(f.tau2_L2, 3), "τ²(L3)": _f(f.tau2_L3, 3),
            "Variance: sampling / L2 / L3 (%)": f"{_f(ps, 1)} / {_f(p2, 1)} / {_f(p3, 1)}",
        })
    note = ("Three-level random-effects model (metafor::rma.mv, random = ~1 | Study/effect, REML). CR2: "
            "clubSandwich cluster-robust inference with Satterthwaite df; model: rma.mv t-test (df = k − 1). "
            "I² = 100 × (1 − sampling-variance share). CR2 df < 4 indicates unreliable robust inference.")
    return SuppTable("S2", "Pooled estimates from three-level random-effects models", pd.DataFrame(rows), note)


def table_s3(results: list[dict]) -> SuppTable:
    rows = []
    for r in results:
        for col in M.candidate_group_columns(r):
            sg = M.subgroup_analysis_3level(r, col)
            q, dfq, pq = sg["qm"]
            for _, t in sg["table"].iterrows():
                rows.append({"Outcome": r["outcome"], "Grouping": col, "Subgroup": t[col], "k": int(t["k"]),
                             "Studies": int(t["n_studies"]), "Hedges' g": _f(t["estimate"]),
                             "95% CI": _ci(t["ci_lb"], t["ci_ub"]), "p": _p(t["pval"]),
                             "Q_M (df)": f"{_f(q)} ({int(dfq)})", "p (Q_M)": _p(pq)})
    note = ("Run only when ≥ 2 categories each contained ≥ 3 independent studies (as in the R pipeline). Subgroup "
            "means from the no-intercept three-level model; Q_M = Wald omnibus test of the moderator (z-based).")
    return SuppTable("S3", "Subgroup analyses", pd.DataFrame(rows), note)


def table_s4(results: list[dict]) -> SuppTable:
    rows = []
    for r in results:
        sd = r["study_df"]
        if len(sd) < 3:
            continue
        po = r["pooled"]
        for t in M.loo_table(sd).itertuples():
            rows.append({"Outcome": r["outcome"], "Omitted study": t.Study, "Hedges' g": _f(t.estimate),
                         "95% CI": _ci(t.ci_lb, t.ci_ub), "p": _p(t.pval), "Δg vs. all studies": _f(t.estimate - po.beta),
                         "τ²": _f(t.tau2, 3), "I² (%)": _f(t.I2, 1)})
    note = ("Study-level effects aggregated within study (compound symmetry, ρ = 0.6) and re-pooled with "
            "rma(REML, test = 'knha') after omitting one study at a time (metafor::leave1out).")
    return SuppTable("S4", "Leave-one-study-out sensitivity analysis", pd.DataFrame(rows), note)


def table_s5(results: list[dict]) -> SuppTable:
    rows = []
    for r in results:
        sd = r["study_df"]
        if len(sd) < 3:
            continue
        for t in M.influence_table(sd).itertuples():
            rows.append({"Outcome": r["outcome"], "Study": t.Study, "Studentized residual": _f(t.rstudent),
                         "DFFITS": _f(t.dffits), "Cook's D": _f(t.cook_d, 3), "Hat": _f(t.hat, 3),
                         "DFBETAS": _f(t.dfbetas), "Cov. ratio": _f(t.cov_r),
                         "Bonferroni outlier": "Yes" if t.formal_outlier_bonferroni else "",
                         "Influential (metafor)": "Yes" if t.influential_metafor else ""})
    note = ("metafor::influence() on the study-level model. Influential: |DFFITS| > 3√(p/(k−p)), Cook's D > χ²₀.₅(p), "
            "hat > 3p/k or |DFBETAS| > 1. Formal outlier: |externally studentized residual| > z(1 − 0.05/2k). "
            "|residual| > 1.96 was used only as a screening reference.")
    return SuppTable("S5", "Influence and outlier diagnostics", pd.DataFrame(rows), note)


def table_s6(results: list[dict], ci_mode: str = "CR2") -> SuppTable:
    rows = []
    for r in results:
        for t in M.robustness_table(r, ci_mode).itertuples():
            rows.append({"Outcome": r["outcome"], "Analysis": t.Analysis, "k": t.k, "Studies": t.Studies,
                         "Hedges' g": _f(t.g), "95% CI": _ci(t.ci_lb, t.ci_ub), "p": _p(t.p),
                         "Removed": t.Removed, "Model / note": t.Note})
    note = ("Outlier exclusion refits the study-level model; influential-study exclusion refits the primary "
            "three-level model (as in the R pipeline). No diagnostic removed studies from the primary analysis.")
    return SuppTable("S6", "Robustness of pooled estimates", pd.DataFrame(rows), note)


def table_s7(results: list[dict]) -> SuppTable:
    rows = []
    for r in results:
        sd = r["study_df"]
        eg, tf = r.get("egger"), r.get("trimfill")
        row = {"Outcome": r["outcome"], "Studies": len(sd)}
        if eg is not None and np.isfinite(eg.p_value):
            row.update({"Egger t (sei)": _f(eg.t_value), "Egger p": _p(eg.p_value)})
        if tf is not None:
            o, a = tf["original"], tf["adjusted"]
            row.update({"Trim-and-fill side": tf["side"], "Imputed (k0)": int(tf["n_missing"]),
                        "Original g (95% CI)": f"{_f(o.beta)} ({_ci(o.ci[0], o.ci[1])})",
                        "Adjusted g (95% CI)": f"{_f(a.beta)} ({_ci(a.ci[0], a.ci[1])})", "Adjusted p": _p(a.p_value)})
        row["Note"] = "k < 10: exploratory" if len(sd) < 10 else ""
        rows.append(row)
    note = ("Egger regression test: metafor::regtest(model = 'rma', predictor = 'sei') on study-level effects "
            "(REML + HKSJ). Trim-and-fill: L0 estimator, side chosen by the sign of the sei slope; the adjusted model "
            "was refitted with imputed studies (z-based CI). Tests with < 10 studies have low power.")
    return SuppTable("S7", "Small-study effects and publication bias", pd.DataFrame(rows), note)


def table_s8(adv: pd.DataFrame | None) -> SuppTable | None:
    if adv is None or len(adv) == 0:
        return None
    a = adv.copy()
    cols = [c for c in ["Outcome", "Analysis", "Status", "Estimate", "95% CI", "p (분류 기준)", "기준",
                        "p (모델 기반)", "Detail"] if c in a.columns]
    a = a[cols].rename(columns={"p (분류 기준)": "p (primary)", "기준": "p basis", "p (모델 기반)": "p (model)"})
    for c in ("Estimate",):
        if c in a.columns:
            a[c] = a[c].map(lambda v: _f(v, 3))
    for c in ("p (primary)", "p (model)"):
        if c in a.columns:
            a[c] = a[c].map(_p)
    note = ("Meta-regression and dose–response: three-level models with CR2 inference; PET-PEESE and the "
            "Vevea–Hedges selection model were run only with ≥ 10 independent studies. SKIPPED = data insufficient.")
    return SuppTable("S8", "Meta-regression, dose–response and advanced bias analyses", a, note)


def table_s9(qc: pd.DataFrame | None) -> SuppTable | None:
    if qc is None or len(qc) == 0:
        return None
    return SuppTable("S9", "Data extraction quality control", qc.copy(),
                     "Rows excluded before analysis and the reason (e.g., SD ≤ 0, missing n).")


def build_tables(results: list[dict], ci_mode: str = "CR2", qc: pd.DataFrame | None = None,
                 adv: pd.DataFrame | None = None, include: tuple[str, ...] | None = None,
                 label_mode: str = "plain") -> list[SuppTable]:
    makers = {
        "S1": lambda: table_s1(results, label_mode), "S2": lambda: table_s2(results),
        "S3": lambda: table_s3(results), "S4": lambda: table_s4(results), "S5": lambda: table_s5(results),
        "S6": lambda: table_s6(results, ci_mode), "S7": lambda: table_s7(results),
        "S8": lambda: table_s8(adv), "S9": lambda: table_s9(qc),
    }
    out = []
    for code, fn in makers.items():
        if include is not None and code not in include:
            continue
        t = fn()
        if t is not None and len(t.df):
            out.append(t)
    return out


def to_xlsx(tables: list[SuppTable]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Contents"
    ws0["A1"] = "Supplementary Tables"
    ws0["A1"].font = Font(bold=True, size=14, color="0F1F3D")
    for i, t in enumerate(tables, start=3):
        ws0.cell(row=i, column=1, value=f"Table {t.code}").font = Font(bold=True)
        ws0.cell(row=i, column=2, value=t.title)
    ws0.cell(row=len(tables) + 4, column=1, value="Abbreviations").font = Font(bold=True)
    ws0.cell(row=len(tables) + 5, column=1, value=ABBR).alignment = Alignment(wrap_text=True, vertical="top")
    ws0.merge_cells(start_row=len(tables) + 5, start_column=1, end_row=len(tables) + 5, end_column=6)
    ws0.row_dimensions[len(tables) + 5].height = 60
    ws0.column_dimensions["A"].width = 14
    ws0.column_dimensions["B"].width = 70

    thin = Side(style="thin", color="3A4A66")
    head_fill = PatternFill("solid", fgColor="EAF1FB")
    for t in tables:
        ws = wb.create_sheet(f"Table {t.code}")
        ncol = max(len(t.df.columns), 1)
        ws.cell(row=1, column=1, value=f"Table {t.code}. {t.title}").font = Font(bold=True, size=12, color="0F1F3D")
        for j, c in enumerate(t.df.columns, start=1):
            cell = ws.cell(row=3, column=j, value=str(c))
            cell.font = Font(bold=True, color="0F1F3D")
            cell.fill = head_fill
            cell.border = Border(top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        prev_outcome = None
        for i, row in enumerate(t.df.itertuples(index=False), start=4):
            for j, v in enumerate(row, start=1):
                if isinstance(v, float) and not np.isfinite(v):
                    v = ""
                if isinstance(v, (np.integer,)):
                    v = int(v)
                if isinstance(v, (np.floating,)):
                    v = float(v)
                cell = ws.cell(row=i, column=j, value=v)
                cell.alignment = Alignment(horizontal="left" if j <= 2 else "center", vertical="center")
            first = row[0]
            if "Outcome" in t.df.columns and prev_outcome is not None and first != prev_outcome:
                for j in range(1, ncol + 1):
                    ws.cell(row=i, column=j).border = Border(top=Side(style="hair", color="C9D1DC"))
            prev_outcome = first
        last = 3 + len(t.df)
        for j in range(1, ncol + 1):
            c = ws.cell(row=last, column=j)
            c.border = Border(bottom=thin, top=c.border.top)
        note_row = last + 2
        ws.cell(row=note_row, column=1, value=f"Note. {t.note}" if t.note else "").alignment = Alignment(
            wrap_text=True, vertical="top")
        ws.cell(row=note_row + 1, column=1, value=ABBR).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=ncol)
        ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=ncol)
        ws.row_dimensions[note_row].height = 48
        ws.row_dimensions[note_row + 1].height = 48
        for j, c in enumerate(t.df.columns, start=1):
            width = max([len(str(c))] + [len(str(v)) for v in t.df.iloc[:, j - 1].head(300)])
            ws.column_dimensions[get_column_letter(j)].width = min(max(9, width + 2), 46)
        ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_docx(tables: list[SuppTable]) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    usable_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000

    def _border(cell, top=None, bottom=None):
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        for edge, val in (("top", top), ("bottom", bottom)):
            if val:
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(val))
                el.set(qn("w:color"), "3A4A66")
                b.append(el)
        tcPr.append(b)

    h = doc.add_paragraph()
    r = h.add_run("Supplementary Tables")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x0F, 0x1F, 0x3D)
    for i, t in enumerate(tables):
        if i:
            doc.add_page_break()
        p = doc.add_paragraph()
        rr = p.add_run(f"Table {t.code}. ")
        rr.bold = True
        p.add_run(t.title).bold = True
        df = t.df.fillna("")
        tbl = doc.add_table(rows=1, cols=len(df.columns))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(df.columns):
            cell = tbl.rows[0].cells[j]
            cell.text = str(c)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(8.5)
            _border(cell, top=8, bottom=6)
        for row in df.itertuples(index=False):
            cells = tbl.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = "" if (isinstance(v, float) and not np.isfinite(v)) else str(v)
                for run in cells[j].paragraphs[0].runs:
                    run.font.size = Pt(8)
        for cell in tbl.rows[-1].cells:
            _border(cell, bottom=8)
        # 열 폭: 내용 길이에 비례(상한), 머리글 행은 페이지마다 반복
        req = []
        for j, c in enumerate(df.columns):
            vals = [len(str(v)) for v in df.iloc[:, j].head(300)]
            longest_word = max([len(w) for w in str(c).split()] + [1])
            req.append(0.17 * max([longest_word] + vals) + 0.35)
        if sum(req) <= usable_cm:
            lens = [r_ * usable_cm / sum(req) for r_ in req]
        else:
            fixed = sum(r_ for r_ in req if r_ <= 3.0)
            flex = [r_ for r_ in req if r_ > 3.0]
            room = max(usable_cm - fixed, 2.2 * max(len(flex), 1))
            scale = room / sum(flex) if flex else 1.0
            lens = [r_ if r_ <= 3.0 else max(2.2, r_ * scale) for r_ in req]
        tbl.autofit = False
        grid = tbl._tbl.tblGrid
        for j, ln in enumerate(lens):
            wcm = Cm(ln * usable_cm / max(sum(lens), usable_cm) if sum(lens) > usable_cm else ln)
            tbl.columns[j].width = wcm
            if j < len(grid.gridCol_lst):
                grid.gridCol_lst[j].w = wcm
            for row in tbl.rows:
                row.cells[j].width = wcm
        trPr = tbl.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)
        if t.note:
            n = doc.add_paragraph()
            n.add_run("Note. ").italic = True
            n.add_run(t.note).font.size = Pt(8)
        a = doc.add_paragraph()
        a.add_run(ABBR).font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# Table S2 (사용자 Supplementary 형식) — 실험군·대조군 원자료(n, Mean, SD)
#   * forest plot과 같은 outcome·같은 행 순서(연도 → 연구 → 중재 → 용량)
#   * Times New Roman 11 pt, 3선 표(머리 위 굵은 선 · 각 행 얇은 선 · 마지막 행 굵은 선),
#     Experimental / Control 병합 머리글, 각주 9 pt(위첨자 번호), A4 가로
#   * 약어는 처음 나오는 outcome 칸에만 위첨자 번호를 달고 각주에 정의
# ===========================================================================
KNOWN_ABBR = {
    "TG": "triglyceride", "TC": "total cholesterol", "WAT": "white adipose tissue", "BAT": "brown adipose tissue",
    "PPARα": "peroxisome proliferator-activated receptor alpha", "PPARγ": "peroxisome proliferator-activated receptor gamma",
    "FAS": "fatty acid synthase", "UCP1": "uncoupling protein 1", "HDL": "high-density lipoprotein",
    "LDL": "low-density lipoprotein", "HDL-C": "high-density lipoprotein cholesterol",
    "LDL-C": "low-density lipoprotein cholesterol", "FFA": "free fatty acid", "NEFA": "non-esterified fatty acid",
    "ALT": "alanine aminotransferase", "AST": "aspartate aminotransferase", "BMD": "bone mineral density",
    "BV/TV": "bone volume fraction", "CSA": "cross-sectional area", "SREBP-1c": "sterol regulatory element-binding protein 1c",
    "ACC": "acetyl-CoA carboxylase", "CPT1": "carnitine palmitoyltransferase 1", "AMPK": "AMP-activated protein kinase",
    "PGC-1α": "peroxisome proliferator-activated receptor gamma coactivator 1-alpha", "SOD": "superoxide dismutase",
    "MDA": "malondialdehyde", "GPx": "glutathione peroxidase", "IL-6": "interleukin 6", "TNF-α": "tumor necrosis factor alpha",
}
S2_CAPTION = "Numerical data for experimental and control groups included in the meta-analysis."
S2_CELL_MAR = {"top": 0, "bottom": 0, "left": 40, "right": 40}   # dxa — 셀 여백(좁게). Word·Google 문서 모두 적용
S2_FOOT_PT = 9
S2_WIDTHS = (1423, 1449, 2197, 1423, 1423, 1423, 1423, 1423, 1424)   # dxa (사용자 원본과 동일)


def fmt_raw(v) -> str:
    """원자료 표기: 정수면 '58', 아니면 입력값 그대로(유효숫자 10자리, 뒤 0 제거)."""
    try:
        x = float(v)
    except (TypeError, ValueError):
        return "" if v is None else str(v)
    if not np.isfinite(x):
        return ""
    s = f"{x:.10g}"
    return s if "e" not in s else f"{x:.6f}".rstrip("0").rstrip(".")


def fmt_n(v) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return "" if v is None else str(v)
    if not np.isfinite(x):
        return ""
    return str(int(round(x))) if abs(x - round(x)) < 1e-9 else f"{x:.10g}"


def fmt_2(v) -> str:
    """Mean·SD 표기: 소수 둘째 자리 고정(예: 200.00, 0.04)."""
    try:
        x = float(v)
    except (TypeError, ValueError):
        return "" if v is None else str(v)
    return f"{x:.2f}" if np.isfinite(x) else ""


def abbr_tokens(name: str, defs: dict) -> list[str]:
    """이름 안의 약어 후보(정의 사전에 있거나 대문자 2개 이상인 토큰, 예: CSA, EPA-PL, AA-Sev)."""
    import re as _re
    toks = _re.findall(r"[A-Za-z0-9α-ωΑ-Ω/\-]+", str(name))
    out = []
    for t in toks:
        t = t.strip("-/")
        if t and (t in defs or sum(ch.isupper() for ch in t) >= 2):
            out.append(t)
    return out


def s2_outcome_name(outcome: str, names: dict | None = None) -> str:
    import re as _re
    if names and str(names.get(outcome, "")).strip():
        return str(names[outcome]).strip()
    return _re.sub(r"\s+", " ", str(outcome).replace("_", " ")).strip()


def _s2_source(results, outcomes, names, study_names):
    """(outcome 표기, 연구 표기, 중재, 원자료 행) 순서대로 — forest plot 행 순서와 같다."""
    by = {r["outcome"]: r for r in results}
    need = {"Mean_treat", "SD_treat", "N_treat", "Mean_control", "SD_control", "N_control"}
    sn = study_names or {}
    for o in outcomes:
        r = by.get(o)
        if r is None or not need.issubset(r["data"].columns):
            continue
        d = r["data"]
        oname = s2_outcome_name(o, names)
        for x in d.itertuples(index=False):
            st_ = " ".join(str(getattr(x, "Study")).split())
            iv = " ".join(str(getattr(x, "Intervention", "")).split()) if "Intervention" in d.columns else ""
            yield oname, sn.get(st_, st_), iv, x


def s2_abbr_candidates(results, outcomes, names=None) -> list[str]:
    """Outcome·Intervention 칸에서 처음 나오는 순서대로 약어 후보."""
    out: list[str] = []
    for oname, _s, iv, _x in _s2_source(results, outcomes, names, None):
        for t in abbr_tokens(oname, KNOWN_ABBR) + abbr_tokens(iv, KNOWN_ABBR):
            if t not in out:
                out.append(t)
    return out


def table_s2_rows(results: list[dict], outcomes: list[str], defs: dict | None = None,
                  names: dict | None = None, study_names: dict | None = None):
    """(rows, footnotes, missing_defs).
    rows: [outcome, outcome_marker, study, intervention, n1, m1, sd1, n2, m2, sd2, intervention_marker].
    약어는 표에서 처음 나오는 칸(Outcome 또는 Intervention)에만 위첨자 번호를 단다. 정의가 빈 약어는 표시하지 않는다."""
    defs = {**KNOWN_ABBR, **(defs or {})}
    defs = {k: v for k, v in defs.items() if str(v).strip()}
    notes = [("n", "number"), ("SD", "standard deviation")]
    seen: set[str] = set()
    missing: list[str] = []

    def marks(text):
        m = []
        for t in abbr_tokens(text, defs):
            if t in seen:
                continue
            seen.add(t)
            if t in defs:
                notes.append((t, defs[t]))
                m.append(f"{len(notes)})")
            elif t not in missing:
                missing.append(t)
        return ",".join(m)

    rows = []
    for oname, study, iv, x in _s2_source(results, outcomes, names, study_names):
        om = marks(oname)
        im = marks(iv)
        rows.append([oname, om, study, iv,
                     fmt_n(getattr(x, "N_treat")), fmt_2(getattr(x, "Mean_treat")), fmt_2(getattr(x, "SD_treat")),
                     fmt_n(getattr(x, "N_control")), fmt_2(getattr(x, "Mean_control")), fmt_2(getattr(x, "SD_control")),
                     im])
    return rows, notes, missing


def _s2_table_element(rows):
    """사용자 원본 Table S2와 같은 서식의 <w:tbl> 요소를 만든다."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    tbl = doc.add_table(rows=2 + len(rows), cols=9)
    t = tbl._tbl
    tblPr = t.tblPr
    for tag, attrs in (("w:tblW", {"w:w": str(sum(S2_WIDTHS)), "w:type": "dxa"}), ("w:jc", {"w:val": "center"}),
                       ("w:tblLayout", {"w:type": "fixed"})):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        tblPr.append(el)
    sty = tblPr.find(qn("w:tblStyle"))
    if sty is not None:
        tblPr.remove(sty)
    tmar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(S2_CELL_MAR[side]))
        e.set(qn("w:type"), "dxa")
        tmar.append(e)
    tblPr.append(tmar)
    grid = t.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), S2_WIDTHS):
        gc.set(qn("w:w"), str(w))

    def cell_fmt(cell, w, top=None, bottom=None, align="center"):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = tcPr.find(qn("w:tcW"))
        if tcW is None:
            tcW = OxmlElement("w:tcW")
            tcPr.append(tcW)
        tcW.set(qn("w:w"), str(w))
        tcW.set(qn("w:type"), "dxa")
        b = OxmlElement("w:tcBorders")
        for edge, sz in (("top", top), ("bottom", bottom)):
            if sz:
                e = OxmlElement(f"w:{edge}")
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), "000000")
                b.append(e)
        tcPr.append(b)
        mar = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(S2_CELL_MAR[side]))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tcPr.append(mar)
        va = OxmlElement("w:vAlign")
        va.set(qn("w:val"), "center")
        tcPr.append(va)
        p = cell.paragraphs[0]
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
        return p

    def run(p, text, bold=False, sup=False):
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(11)
        r.font.name = "Times New Roman"
        rPr = r._r.get_or_add_rPr()
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.insert(0, rf)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), "Times New Roman")
        if sup:
            r.font.superscript = True
        return r

    h0, h1 = tbl.rows[0].cells, tbl.rows[1].cells
    for j in range(3):                                   # Outcome / Study / Intervention: 세로 병합
        p = cell_fmt(h0[j], S2_WIDTHS[j], top=12)
        run(p, ["Outcome", "Study", "Intervention"][j], bold=True)
        cell_fmt(h1[j], S2_WIDTHS[j], bottom=4)
        h0[j].merge(h1[j])
    for start, label in ((3, "Experimental"), (6, "Control")):  # 그룹 머리글: 가로 병합
        m = h0[start].merge(h0[start + 2])
        p = cell_fmt(m, sum(S2_WIDTHS[start:start + 3]), top=12, bottom=6)
        for extra in m.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        run(p, label, bold=True)
    for j, (lab, sup) in enumerate([("n", "1)"), ("Mean", ""), ("SD", "2)")] * 2):
        c = tbl.rows[1].cells[3 + j]
        p = cell_fmt(c, S2_WIDTHS[3 + j], top=6, bottom=4)
        run(p, lab, bold=True)
        if sup and j < 3:
            run(p, sup, bold=True, sup=True)
    for tr in (tbl.rows[0], tbl.rows[1]):
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))
    last = len(rows) - 1
    for i, rw in enumerate(rows):
        cells = tbl.rows[2 + i].cells
        vals = [rw[0], rw[2], rw[3]] + rw[4:10]
        for j, v in enumerate(vals):
            p = cell_fmt(cells[j], S2_WIDTHS[j], top=4, bottom=12 if i == last else 4,
                         align="left" if j < 3 else "center")
            run(p, v)
            if j == 0 and rw[1]:
                run(p, rw[1], bold=True, sup=True)
            if j == 2 and len(rw) > 10 and rw[10]:
                run(p, rw[10], bold=True, sup=True)
    for tr in tbl.rows:
        trPr = tr._tr.get_or_add_trPr()
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        trPr.append(jc)
    return t


def _s2_footnote_runs(p, notes):
    """각주 문단: 모든 글자(위첨자 번호 포함)와 문단 기호까지 9 pt Times New Roman, 줄 간격 1.0."""
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def size(rPr):
        for tag in ("w:sz", "w:szCs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rPr.append(el)
            el.set(qn("w:val"), str(S2_FOOT_PT * 2))

    def fonts(rPr):
        rf = rPr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rPr.insert(0, rf)
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), "Times New Roman")

    def r(text, sup=False):
        x = p.add_run(text)
        x.font.size = Pt(S2_FOOT_PT)
        x.font.name = "Times New Roman"
        rPr = x._r.get_or_add_rPr()
        fonts(rPr)
        size(rPr)
        x.font.superscript = sup
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pPr = p._p.get_or_add_pPr()
    mark = pPr.find(qn("w:rPr"))
    if mark is None:
        mark = OxmlElement("w:rPr")
        pPr.append(mark)
    fonts(mark)
    size(mark)
    for i, (ab, de) in enumerate(notes, start=1):
        r(f"{i})", sup=True)
        r(f"{ab}: {de}")
        r("; " if i < len(notes) else ".")


def build_table_s2_docx(results: list[dict], outcomes: list[str], defs: dict | None = None,
                        caption: str = S2_CAPTION, names: dict | None = None,
                        study_names: dict | None = None) -> tuple[bytes, dict]:
    """Table S2만 들어 있는 Word 파일(A4 가로, 사용자 원본과 같은 여백·서식)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Pt, Twips

    rows, notes, missing = table_s2_rows(results, outcomes, defs, names, study_names)
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(16838), Twips(11906)
    sec.top_margin, sec.bottom_margin = Twips(1077), Twips(1077)
    sec.left_margin, sec.right_margin = Twips(1440), Twips(1701)
    cap = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_after = Pt(4)
    for text, bold in (("Table S2. ", True), (caption, False)):
        rr = cap.add_run(text)
        rr.bold = bold
        rr.font.size = Pt(11)
        rr.font.name = "Times New Roman"
    tbl_el = _s2_table_element(rows)
    cap._p.addnext(tbl_el)
    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(2)
    foot.paragraph_format.space_after = Pt(0)
    _s2_footnote_runs(foot, notes)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"rows": len(rows), "notes": notes, "missing_defs": missing}


def insert_table_s2(docx_bytes: bytes, results: list[dict], outcomes: list[str], defs: dict | None = None):
    """사용자 Supplementary 워드에서 'Table S2.' 캡션 다음 표와 각주를 새 값으로 바꾼다.
    다른 표·그림·문단·구역(sectPr)은 그대로 둔다. 반환: (새 docx bytes, 정보 dict)."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(docx_bytes))
    body = doc.element.body
    cap = None
    for el in body.iterchildren():
        if el.tag == qn("w:p") and Paragraph(el, doc).text.strip().startswith("Table S2"):
            cap = el
            break
    if cap is None:
        raise ValueError("워드 파일에서 'Table S2.'로 시작하는 캡션을 찾지 못했습니다.")
    old_tbl = cap.getnext()
    while old_tbl is not None and old_tbl.tag not in (qn("w:tbl"), qn("w:p")):
        old_tbl = old_tbl.getnext()
    if old_tbl is None or old_tbl.tag != qn("w:tbl"):
        raise ValueError("'Table S2.' 캡션 바로 다음에 표가 없습니다.")
    old_rows = []
    from docx.table import Table
    for tr in Table(old_tbl, doc).rows[2:]:
        old_rows.append([c.text.strip() for c in tr.cells])
    rows, notes, missing = table_s2_rows(results, outcomes, defs)
    new_tbl = _s2_table_element(rows)
    old_tbl.addprevious(new_tbl)
    foot = old_tbl.getnext()
    body.remove(old_tbl)
    if foot is not None and foot.tag == qn("w:p"):
        p = Paragraph(foot, doc)
        txt = p.text.strip()
        if txt[:2] in ("1)", "1 ") or txt.startswith("1"):
            for r in list(foot.findall(qn("w:r"))):
                foot.remove(r)
            _s2_footnote_runs(p, notes)
    import re as _re
    norm = lambda s: _re.sub(r"\d\)$", "", s)  # noqa: E731
    diffs = []
    for i in range(max(len(old_rows), len(rows))):
        a = old_rows[i] if i < len(old_rows) else None
        b = rows[i] if i < len(rows) else None
        av = None if a is None else [norm(a[0])] + a[1:]
        bv = None if b is None else [b[0], b[2], b[3]] + b[4:10]
        if av != bv:
            diffs.append({"row": i + 1, "기존": " | ".join(av) if av else "(없음)", "새 표": " | ".join(bv) if bv else "(없음)"})
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"rows": len(rows), "old_rows": len(old_rows), "diffs": diffs, "notes": notes,
                            "missing_defs": missing}
