"""SR Studio V36 자체 점검.  실행: python selftest_v36.py

A. 스크리닝 안전 수정
   - 기본 규칙이 비어 있는가 / 라벨 0개에서 규칙이 자동 제외를 하지 않는가(V35 구멍)
   - 같은 코퍼스를 다른 행 순서로 넣어도 결과·지문이 같은가
   - 감사(audit)에서 개발용 라벨이 '읽은 편수'에 들어가지 않는가
   - 규칙 입력 파서, 실행 manifest
B. 메타분석 수치 = R(metafor/clubSandwich) 출력 (인천대 카로티노이드 TG, r_outputs)
   - 3-level pooled(모델·CR2·PI), leave1out, influence(rstudent·Cook), baujat, Egger, trim-and-fill, subgroup Q_M
C. Figure·표 생성
   - forest V1/V2, subgroup, sensitivity 5종, trim-and-fill 3종, 4개 형식 저장
   - Supplementary xlsx/docx, 섹션 zip
"""
from __future__ import annotations

import io
import json
import sys
import warnings
import zipfile

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

import screening as S  # noqa: E402

# ---------------------------------------------------------------------------
# R 참조값 (metafor 4.x / clubSandwich, 01_stat_analysis.R 실행 결과 — TG)
# ---------------------------------------------------------------------------
EFF = [('Woo (2010)', 'Fucoxanthin', -1.437503217549, 0.251660387512), ('Woo (2010)', 'Fucoxanthin', -1.194713675736, 0.235683519175),
       ('Jia (2016)', 'Astaxanthin', -1.027441157425, 0.226390883299), ('Jia (2016)', 'Astaxanthin', -1.531182502371, 0.258612996389),
       ('Koo (2019)', 'Fucoxanthin', -0.409643705207, 0.20419519913), ('Wang (2019)', 'Astaxanthin', -1.173394605744, 0.468842745039),
       ('Wang (2019)', 'Astaxanthin', -1.508859252518, 0.513832812195), ('Sun (2020)', 'Fucoxanthin', -2.406309708835, 0.430947700464),
       ('Sun (2020)', 'Fucoxanthin', -3.886040949464, 0.721916070653), ('Joo (2021)', 'Capsanthin', -2.606197554055, 0.57076479763),
       ('Boshra (2022)', 'Astaxanthin', -0.90578314523, 0.220511077655), ('Shatoor (2022)', 'Astaxanthin', -5.364394037143, 1.149272605804),
       ('Gopal (2022)', 'Lutein', -0.494206803579, 0.343510015196), ('Gopal (2022)', 'Lutein', -0.853814891041, 0.36370832784),
       ('Hao (2024)', 'Fucoxanthin', -1.883514687074, 0.320767432678), ('Hao (2024)', 'Fucoxanthin', -2.237918084853, 0.361341037625),
       ('Hao (2024)', 'Fucoxanthin', -2.405129365071, 0.382906868409), ('Ding (2025)', 'Fucoxanthin', -0.318291571099, 0.405065476212),
       ('Qiu (2015)', 'Lutein', 0.644795650321, 0.262992544709), ('Qiu (2015)', 'Lutein', 0.78030790291, 0.269027513229),
       ('Qiu (2015)', 'Lutein', 0.644795650321, 0.262992544709)]
SD = [('Woo (2010)', -1.306167080276, 0.194571170811), ('Jia (2016)', -1.237616297669, 0.192507279533),
      ('Koo (2019)', -0.409643705207, 0.20419519913), ('Wang (2019)', -1.321958757355, 0.391630309972),
      ('Sun (2020)', -2.700963082521, 0.411774762073), ('Joo (2021)', -2.606197554055, 0.57076479763),
      ('Boshra (2022)', -0.90578314523, 0.220511077655), ('Shatoor (2022)', -5.364394037143, 1.149272605804),
      ('Gopal (2022)', -0.661180589646, 0.282483741065), ('Hao (2024)', -2.122782487755, 0.257334477049),
      ('Ding (2025)', -0.318291571099, 0.405065476212), ('Qiu (2015)', 0.687757568227, 0.194289279113)]
POOLED = {'mu': -1.4108372217, 'ci_lb': -2.2175805345, 'ci_ub': -0.6040939089, 'cr2_ci_lb': -2.2603455961,
          'cr2_ci_ub': -0.5613288473, 'pi_lb': -4.107622604, 'pi_ub': 1.2859481606, 'tau2_L3': 1.521821663}
RSTUDENT = [0.0616148775, 0.1116776011, 0.7315749691, 0.0464861734, -0.958785193, -0.8376106296, 0.3521666435,
            -3.0231436335, 0.523809069, -0.5390224996, 0.7535527059, 1.8018140885]
COOK = [0.0031756135, 0.0053332953, 0.0647846805, 0.0022302367, 0.0855096839, 0.0571677312, 0.0219031803,
        0.4132030039, 0.0376627223, 0.0237300292, 0.0597177564, 0.1737049769]
LOO = [-1.3991122623, -1.4059189935, -1.4799986788, -1.3953869507, -1.2567588544, -1.2785227549, -1.4365174775,
       -1.1137461441, -1.4553218401, -1.3132372302, -1.4758533833, -1.5462222006]
BJX = [0.0031167254, 0.0122357942, 0.5914397217, 0.001659864, 0.9822888801, 0.7776018509, 0.1386355124,
       6.3011182486, 0.3083555792, 0.34152421, 0.6285824613, 2.7141418471]
BJY = [0.0025754306, 0.0043286523, 0.0558864685, 0.0018293199, 0.0785093665, 0.0515178074, 0.0180368837,
       0.7514298141, 0.0316767073, 0.020039494, 0.05228507, 0.2003166561]
EGGER = (-3.7536710458, 0.0037610575)
TF = (-1.37611167804252, -2.27445462412725, -0.477768731957785, 0)
SG = [('Astaxanthin', -1.9086125888, -3.254593395), ('Fucoxanthin', -1.4503690972, -2.607089272)]
SG_QM = (0.2561086844, 0.6128062341)
TOL = 1e-6


def _close(a, b, what, tol=TOL):
    a, b = np.asarray(a, float), np.asarray(b, float)
    assert a.shape == b.shape and np.allclose(a, b, atol=tol, rtol=0), f"{what}: {a} vs {b}"


def check_screening() -> None:
    S._HAS_SENTENCE_TRANSFORMERS = False
    assert S.GATE_RULES_DEFAULT == [], "기본 규칙은 비어 있어야 함"
    rules = S.parse_gate_rules_text("노출: nitrosamin*, NDMA, heterocyclic amine*\n결과: re:\\bldl\\b\n# 주석")
    assert [n for n, _ in rules] == ["노출 없음", "결과 없음"]
    import re
    assert re.search(rules[0][1], "n-nitrosamines") and not re.search(rules[0][1], "xndmax")
    back = S.parse_gate_rules_text(S.gate_rules_to_text(S.LEGACY_NITROSAMINE_CVD_GATE_RULES))
    assert [n for n, _ in back] == [n for n, _ in S.LEGACY_NITROSAMINE_CVD_GATE_RULES]
    try:
        S.parse_gate_rules_text("bad: re:(unclosed")
        raise AssertionError("잘못된 정규식이 통과됨")
    except ValueError:
        pass

    rng = np.random.default_rng(0)
    filler = ["study", "analysis", "results", "method", "group", "level"]

    def ab(core):
        return core + ". " + " ".join(rng.choice(filler, 30)) + "."
    rows = [(f"Fucoxanthin reduces obesity in mice {i}", ab("serum triglyceride and body weight in obese mice"), 1) for i in range(20)]
    rows += [(f"Lycopene and prostate cancer cell line {i}", ab("cell proliferation assay"), 0) for i in range(400)]
    c = pd.DataFrame(rows, columns=["Title", "Abstract", "truth"]).sample(frac=1, random_state=1).reset_index(drop=True)
    c0 = c[["Title", "Abstract"]].copy()
    c0["Human_Label"] = ""
    r0 = S.rule_only_screen(c0, "P: obese\nI: carotenoid\nO: weight", gate_rules=S.LEGACY_NITROSAMINE_CVD_GATE_RULES)
    assert r0.metrics["auto_excluded_n"] == 0, "라벨 0개인데 규칙이 자동 제외를 함(V35 구멍)"
    r1 = S.rule_only_screen(c0, "P: obese\nI: carotenoid\nO: weight")
    assert r1.metrics["auto_excluded_n"] == 0 and r1.metrics["quality_gate_reasons"]

    n = 260
    y = (rng.random(n) < 0.1).astype(int)
    pos, neg = ["nitrosamine cardiac rat", "heart aorta mice"], ["cell line assay", "human cohort diet"]
    t, a = [], []
    for k in range(n):
        src = pos if (y[k] == 1) != (rng.random() < 0.3) else neg
        t.append(f"{rng.choice(src)} paper {k}")
        a.append(ab(rng.choice(src)))
    d = pd.DataFrame({"Title": t, "Abstract": a, "Human_Label": y})
    crit = "P: rat\nI: nitrosamine\nO: cardiovascular"
    ra = S.train_and_predict(d, criteria_text=crit)
    rb = S.train_and_predict(d.sample(frac=1, random_state=7).reset_index(drop=True), criteria_text=crit)
    pa = ra.predictions.set_index("Title")["AI_Probability"]
    pb = rb.predictions.set_index("Title")["AI_Probability"].loc[pa.index]
    assert np.array_equal(pa.to_numpy(), pb.to_numpy()), "행 순서에 따라 결과가 달라짐"
    assert ra.metrics["run_fingerprint"] == rb.metrics["run_fingerprint"], "행 순서에 따라 지문이 달라짐"
    rc = S.train_and_predict(d, criteria_text=crit + " extra")
    assert rc.metrics["run_fingerprint"] != ra.metrics["run_fingerprint"], "PICO가 지문에 반영되지 않음"
    man = json.loads(S.build_run_manifest(ra, crit, "", [], d).decode("utf-8"))
    for key in ("run_fingerprint", "pico_sha256", "software", "output_sha256", "input_corpus_sha256"):
        assert man.get(key), f"manifest에 {key} 없음"

    pool = ra.predictions.copy()
    pool["AI_Recommendation"] = "안전 제외 후보"
    st = S.audit_risk_strata(pool)
    assert int(st["읽은_편수"].sum()) == 0, "개발 라벨이 감사 근거로 집계됨"
    assert "개발라벨_편수(참고·제외)" in st.columns


def _tg_result(ci_mode="model"):
    import meta_sections as M
    eff = pd.DataFrame(EFF, columns=["study", "Intervention", "yi", "vi"])
    return M.result_from_effects(eff, "TG", ci_mode)


def check_meta_numbers() -> None:
    import meta_sections as M
    r = _tg_result("model")
    f = r["fit"]
    _close([f.mu, f.ci_lb, f.ci_ub, f.cr2_ci_lb, f.cr2_ci_ub, f.pi_lb, f.pi_ub, f.tau2_L3],
           [POOLED[k] for k in ("mu", "ci_lb", "ci_ub", "cr2_ci_lb", "cr2_ci_ub", "pi_lb", "pi_ub", "tau2_L3")],
           "3-level pooled", tol=1e-5)
    sd = r["study_df"].set_index("study").loc[[s for s, _, _ in SD]].reset_index()
    _close(sd["yi"], [y for _, y, _ in SD], "CS 집계 yi")
    inf = M.influence_table(sd)
    _close(inf["rstudent"], RSTUDENT, "rstudent")
    _close(inf["cook_d"], COOK, "Cook's D")
    _close(M.loo_table(sd)["estimate"], LOO, "leave1out")
    bj = M.baujat_table(sd)
    _close(bj["x_heterogeneity"], BJX, "Baujat x")
    _close(bj["y_influence"], BJY, "Baujat y")
    eg = r["egger"]
    _close([eg.t_value, eg.p_value], EGGER, "Egger")
    tf = r["trimfill"]
    _close([tf["adjusted"].beta, tf["adjusted"].ci[0], tf["adjusted"].ci[1], tf["n_missing"]], TF, "trim-and-fill")
    sg = M.subgroup_analysis_3level(r, "Intervention")
    t = sg["table"].set_index("Intervention")
    for name, est, lb in SG:
        _close([t.loc[name, "estimate"], t.loc[name, "ci_lb"]], [est, lb], f"subgroup {name}")
    _close([sg["qm"][0], sg["qm"][2]], SG_QM, "Q_M")
    rob = M.robustness_table(r, "model").set_index("Analysis")
    _close(rob.loc["Excluding Bonferroni outliers", "g"], -1.11374614408972, "outlier-removed (R)")
    gt = M.gosh_table(sd)
    assert len(gt) == 2 ** 12 - 1


def check_figures_tables() -> None:
    import forest_styles as F
    import meta_sections as M
    import supplementary as SUP
    r = _tg_result("CR2")
    st = M.default_settings("TG")
    assert st["favours"] == ("Favours Intervention", "Favours Control")
    assert M.default_settings("PPARα")["favours"] == ("Favours Control", "Favours Intervention")
    kinds = ["forest_V1", "forest_V2", "subgroup", "leave1out", "influence", "baujat", "gosh", "robustness",
             "trimfill", "funnel", "trimfill_compare"]
    for k in kinds:
        fig = M.make_figure(k, r, st, "CR2", "Intervention" if k == "subgroup" else None)
        assert fig is not None, f"{k} figure 없음"
        for fmt in ("png", "tiff", "pdf", "svg"):
            b = F.save_figure(fig, fmt, 120)
            assert len(b) > 2000, f"{k}.{fmt} 비어 있음"
    w = M.forest_fig(r, st).get_size_inches()[0]
    assert abs(w - F.FIGW) < 0.05, f"V1 인쇄 폭이 6.3 in가 아님: {w}"
    tabs = SUP.build_tables([r], "CR2", None, None)
    assert {t.code for t in tabs} >= {"S1", "S2", "S3", "S4", "S5", "S6", "S7"}
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(SUP.to_xlsx(tabs)))
    assert "Table S2" in wb.sheetnames
    from docx import Document
    doc = Document(io.BytesIO(SUP.to_docx(tabs)))
    assert len(doc.tables) == len(tabs)
    z = zipfile.ZipFile(io.BytesIO(M.build_section_zip([r], "trimfill", formats=("png",), dpi=80)))
    names = z.namelist()
    assert any(n.startswith("03_TrimFill_Plots/trimfill_TG") for n in names) and "trimfill_numbers.xlsx" in names


def check_auto_gate() -> None:
    """자동 개념 게이트: 코드 수정·정규식 입력 없이 PICO + Include 라벨만으로 블록을 만들고,
    라벨되지 않은 적격 문헌을 떨어뜨리지 않는가(합성 코퍼스)."""
    rng = np.random.default_rng(0)
    filler = ["study", "analysis", "results", "significant", "group", "levels", "week", "treatment"]

    def ab(t):
        return t + ". " + " ".join(rng.choice(filler, 30)) + "."
    exp = ["N-nitrosodiethylamine (NDEA)", "N-nitrosodimethylamine (NDMA)", "heterocyclic amines such as PhIP",
           "nitrosamine", "MeIQx"]
    out = ["aortic atherosclerosis", "serum LDL cholesterol", "myocardial injury", "endothelial dysfunction"]
    ani = ["Wistar rats", "C57BL/6 mice", "Sprague-Dawley rats"]
    rows = [(f"Effect of {rng.choice(exp)} on {rng.choice(out)} in {rng.choice(ani)} {i}",
             ab(f"{rng.choice(ani)} were exposed to {rng.choice(exp)} and {rng.choice(out)} was assessed"), 1)
            for i in range(30)]
    noise = [("NDEA-induced hepatocarcinogenesis in rats", "liver tumor incidence after NDEA in rats"),
             ("Nitrosamine formation in processed meat", "analytical chemistry of nitrosamines in cured meat"),
             ("Statins and atherosclerosis in mice", "atorvastatin reduced aortic plaque in ApoE mice")]
    rows += [(f"{noise[i % 3][0]} {i}", ab(noise[i % 3][1]), 0) for i in range(900)]
    d = pd.DataFrame(rows, columns=["Title", "Abstract", "truth"]).sample(frac=1, random_state=2).reset_index(drop=True)
    lab = list(d.index[d.truth == 1][:8]) + list(d.index[d.truth == 0][:192])
    d["Human_Label"] = np.nan
    d.loc[lab, "Human_Label"] = d.loc[lab, "truth"]
    crit = ("P: rats or mice (animal models)\nI: N-nitrosamines, heterocyclic amines\n"
            "O: cardiovascular outcomes, atherosclerosis, lipid profile")
    g = S.derive_auto_gate(d[["Title", "Abstract", "Human_Label"]], crit)
    assert g["usable"] and g["rules"], f"자동 블록이 만들어지지 않음: {g['reasons']}"
    rr = S.rule_only_screen(d[["Title", "Abstract", "Human_Label"]], crit, gate_rules=g["rules"])
    p = rr.predictions.merge(d[["Title", "truth"]], on="Title")
    lost = int((p["AI_Recommendation"].eq("안전 제외 후보") & p["truth"].eq(1)).sum())
    assert lost == 0, f"자동 게이트가 적격 문헌 {lost}편을 제외함"
    k = S.derive_auto_gate(d[["Title", "Abstract", "Human_Label"]], "P: 비만 동물\nI: 카로티노이드\nO: 체중")
    assert not k["usable"] and k["reasons"], "한국어 PICO에서 근거 없는 블록이 만들어짐"
    terms_all = " ".join(g["blocks"]["용어"])
    assert ("ndea" in terms_all) and ("n-nitrosodiethylamine" in terms_all), "약어·형태 확장 실패"


def check_table_s2() -> None:
    import supplementary as SUP
    from docx import Document
    r = _tg_result("model")
    r["data"]["Mean_treat"], r["data"]["SD_treat"], r["data"]["N_treat"] = 1.03, 0.126, 10
    r["data"]["Mean_control"], r["data"]["SD_control"], r["data"]["N_control"] = 1.33, 0.253, 10
    rows, notes, missing = SUP.table_s2_rows([r], ["TG"])
    assert len(rows) == len(r["data"]) and rows[0][1] == "3)" and rows[1][1] == ""
    assert [a for a, _ in notes] == ["n", "SD", "TG"] and not missing
    assert SUP.fmt_raw(58.0) == "58" and SUP.fmt_raw(0.126) == "0.126"
    b, info = SUP.build_table_s2_docx([r], ["TG"])
    doc = Document(io.BytesIO(b))
    assert doc.paragraphs[0].text.startswith("Table S2.") and len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.rows) == 2 + len(rows) and t.rows[0].cells[3].text == "Experimental"
    b2, info2 = SUP.insert_table_s2(b, [r], ["TG"])
    assert info2["old_rows"] == info2["rows"] and not info2["diffs"], "같은 데이터로 교체했는데 차이가 남"


def check_app_imports() -> None:
    import ast
    for fname, mod in (("app.py", "screening"), ("meta_page.py", None)):
        tree = ast.parse(open(fname, encoding="utf-8").read())
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom) and node.module and (mod is None or node.module == mod):
                m = __import__(node.module)
                missing = [a.name for a in node.names if not hasattr(m, a.name)]
                assert not missing, f"{fname}: {node.module}에 없는 이름 {missing}"
    src = open("app.py", encoding="utf-8").read()
    assert "use_container_width" not in src


if __name__ == "__main__":
    check_app_imports()
    check_screening()
    check_meta_numbers()
    check_figures_tables()
    check_auto_gate()
    check_table_s2()
    print("V36 self-test: PASS")
    sys.exit(0)
